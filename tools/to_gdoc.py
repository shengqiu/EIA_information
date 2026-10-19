#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""给 Drive 上「湖州线索/<项目>/」里的附件各生成一份 Google 文档（在 to_drive.py 之后跑，不动原文件）。

    湖州线索/<项目>/
      ├── 网页内容.txt                       ← to_drive.py（不变）
      ├── 报告表.pdf                          ← to_drive.py 传的原文件；文档版生成成功后移到 Drive 回收站
      └── 报告表（文档版）                    ← 本脚本：Google 文档

2026-10-19 起按 Frank 要求不保留附件原文件：文档版生成成功才删（进回收站，30 天可恢复）；
转换失败的、转不了的（图片、表格等）原文件保留。只影响本脚本处理到的附件，之前的旧原文件不动。

文档版的内容：
- Word（.doc/.docx）：直接让 Google 转换，原文图片自动保留；
- 文字版 PDF：文字层原文（没有错字）+ 图页（平面图 / 流程图 / 示意图等）整页插图；
- 扫描版 PDF：Google OCR 文字（分段 ≤10 页、≤2MB 交给 Google 转换识别，绕开它的 OCR 限制）+ 图页整页插图。
其他文件（图片、表格等）不转。

逻辑与 shengqiu/wuxi-leads 的 wuxi/attach2doc.py、wuxi/leaddoc.py 一致（2026-10-18 在那边用真实附件测过：
157 页文字版、80 页扫描件、367 页报告书都正常）。wuxi-leads 的第 1 步 / 第 2 步会直接读这里的文档版。

断点续跑：data/gdoc_state.json 按源文件记账（gdoc|<OSS 直链>#<压缩包内文件名>），转过的跳过。
默认只处理最近 7 天有公告的项目（--since 0 = 全部）。

用法：
    python3 tools/to_gdoc.py                    # 最近 7 天的项目
    python3 tools/to_gdoc.py --since 30
    python3 tools/to_gdoc.py --project 化工 --dry-run
"""
import argparse
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import datetime, timedelta, timezone

import pymupdf
from docx import Document
from docx.shared import Cm, Pt

from eia_common import (ARCHIVE_EXT, CST, build_projects, download, extract_bsdtar, extract_zip, load_announcements,
                        load_links, log, sanitize, sniff_archive_kind)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATE_PATH = os.path.join(ROOT, "data", "gdoc_state.json")
REMOTE = "gdrive:"
SUFFIX = "（文档版）"
CONVERTIBLE = (".pdf", ".doc", ".docx")
MIME = {".pdf": "application/pdf", ".doc": "application/msword",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}

# ---------------------------------------------------------------- 判断 & 生成（同 wuxi-leads attach2doc.py）

_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
FIG_WORDS = re.compile(r"平面(布置)?图|布置图|流程图|示意图|位置图|四至图|周边(环境|关系)|附图|监测点位|厂区平面|照片|图\s*\d")
FEW_CHARS, CAPTION_MAX = 250, 600
OCR_PAGES, OCR_MAX_BYTES = 10, 1_900_000
OCR_NOTE = ("本文档由扫描件自动生成：文字是 Google OCR 识别结果，仅供检索和快速浏览，中文 OCR 会有错字"
            "（如「氩/氧」「氮/氢」），数字以原件为准；平面图、流程图等图页按原样插入。")


def is_scanned(data, pages=10, min_chars=50):
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
        n = min(len(doc), pages)
        return n == 0 or sum(len(doc[i].get_text().strip()) for i in range(n)) / n < min_chars
    except Exception:
        return True


def _rough_ocr(page):
    if not shutil.which("tesseract"):
        return None
    png = page.get_pixmap(dpi=90, colorspace=pymupdf.csGRAY).tobytes("png")
    p = subprocess.run(["tesseract", "stdin", "stdout", "-l", "chi_sim", "--psm", "3"],
                       input=png, capture_output=True, timeout=120)
    return re.sub(r"\s+", "", p.stdout.decode("utf-8", "ignore"))


def classify(data, scanned):
    doc = pymupdf.open(stream=data, filetype="pdf")
    pages = []
    for i, page in enumerate(doc):
        area = page.rect.width * page.rect.height or 1
        if not scanned:
            text = page.get_text().strip()
            img_area = sum(r.width * r.height for img in page.get_images(full=True)
                           for r in page.get_image_rects(img[0]))
            dense = len(re.sub(r"\s+", "", text))
            figure = dense < FEW_CHARS or img_area / area > 0.5 or len(page.get_drawings()) > 300 \
                or (dense < CAPTION_MAX and bool(FIG_WORDS.search(text)))
        else:
            text, rough = "", _rough_ocr(page)
            figure = True if rough is None else \
                len(rough) < FEW_CHARS or (len(rough) < CAPTION_MAX and bool(FIG_WORDS.search(rough)))
        pages.append({"no": i + 1, "text": text, "figure": figure})
    return pages


def ocr_chunks(data):
    src = pymupdf.open(stream=data, filetype="pdf")
    out, start = [], 0
    while start < len(src):
        n = min(OCR_PAGES, len(src) - start)
        for dpi, quality in ((150, 70), (120, 60), (100, 50), (80, 45)):
            doc = pymupdf.open()
            for i in range(start, start + n):
                page = src[i]
                pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
                new = doc.new_page(width=page.rect.width, height=page.rect.height)
                new.insert_image(new.rect, stream=pix.tobytes("jpeg", jpg_quality=quality))
            buf = doc.tobytes(garbage=3, deflate=True)
            if len(buf) <= OCR_MAX_BYTES:
                break
        if len(buf) > OCR_MAX_BYTES and n > 1:
            n = max(1, n // 2)
            continue
        out.append((start + 1, start + n, buf))
        start += n
    return out


def build_docx(title, data, pages, ocr_parts=None, note=""):
    src = pymupdf.open(stream=data, filetype="pdf")
    d = Document()
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_heading(title, level=1)
    if note:
        d.add_paragraph(note).runs[0].italic = True

    def add_text(text):
        # PDF 文字层里偶尔夹着控制字符，Word（XML）不接受，先去掉（2026-10-19 青岑生物医药的报告表遇到）
        text = _CTRL.sub("", text)
        for line in text.splitlines():
            if line.strip():
                d.add_paragraph(line.strip())

    def add_figure(no):
        d.add_paragraph(f"【第 {no} 页 · 图】")
        d.add_picture(io.BytesIO(src[no - 1].get_pixmap(dpi=130).tobytes("jpeg", jpg_quality=75)), width=Cm(16))

    if ocr_parts is None:
        for p in pages:
            d.add_paragraph(f"—— 第 {p['no']} 页 ——").runs[0].bold = True
            add_text(p["text"])
            if p["figure"]:
                add_figure(p["no"])
    else:
        figs = {p["no"] for p in pages if p["figure"]}
        for a, b, text in ocr_parts:
            d.add_heading(f"第 {a}–{b} 页", level=2)
            add_text(text)
            for no in range(a, b + 1):
                if no in figs:
                    add_figure(no)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- rclone

def rc(*args, timeout=900, check=True, input_bytes=None):
    """跑 rclone；限流 / 瞬时失败重试（同 wuxi-leads gdrive.py：共享客户端常 rateLimitExceeded）。"""
    last = ""
    for attempt in range(3):
        p = subprocess.run(["rclone", *args, "--retries", "3", "--low-level-retries", "5",
                            "--tpslimit", "4", "--tpslimit-burst", "4"],
                           capture_output=True, input=input_bytes, timeout=timeout)
        if p.returncode == 0 or not check:
            return p
        last = (p.stderr.decode("utf-8", "ignore").strip().splitlines() or [f"exit {p.returncode}"])[-1]
        time.sleep(30 * (attempt + 1) if "rateLimit" in last else 5)
    raise RuntimeError(f"rclone {args[0]} 失败：{last[:200]}")


def trash_raw(folder, name):
    """删掉 to_drive.py 传上去的原文件（进 Drive 回收站）。返回是否删掉。"""
    p = rc("deletefile", f"{folder}/{name}", check=False)
    return p.returncode == 0


def upload_as_gdoc(local, ext, dest_folder, name):
    """本地文件 → Google 文档（Drive 导入转换）。"""
    rc("copyto", local, f"{dest_folder}/{name}{ext}", "--drive-import-formats", ext.lstrip("."),
       "--drive-allow-import-name-change")


def google_ocr(data, tmp_local, tmp_remote):
    """扫描版 PDF 分段交给 Google 转换（自动 OCR），返回 [(起页, 止页, 文字)]。"""
    parts = []
    try:
        for a, b, chunk in ocr_chunks(data):
            name = f"p{a:04d}"
            path = os.path.join(tmp_local, name + ".pdf")
            with open(path, "wb") as fh:
                fh.write(chunk)
            upload_as_gdoc(path, ".pdf", tmp_remote, name)
            text = None
            for _ in range(12):           # 刚转换完的文档要等目录列表跟上
                p = rc("cat", f"{tmp_remote}/{name}.txt", "--drive-export-formats", "txt", check=False)
                if p.returncode == 0:
                    text = p.stdout.decode("utf-8", "ignore")
                    break
                time.sleep(5)
            if text is None:
                raise RuntimeError(f"第 {a}-{b} 页 OCR 文字导出失败")
            parts.append((a, b, text))
    finally:
        rc("purge", tmp_remote, "--drive-use-trash=false", check=False)
    return parts


def convert_one(path, name, folder, tmp_local):
    """单个文件 → 「<名称>（文档版）」Google 文档。返回说明文字。"""
    stem, ext = os.path.splitext(name)
    ext = ext.lower()
    title = stem + SUFFIX
    if ext in (".doc", ".docx"):
        upload_as_gdoc(path, ext, folder, title)
        return "Word 原样转换"
    with open(path, "rb") as fh:
        data = fh.read()
    scanned = is_scanned(data)
    pages = classify(data, scanned)
    figs = sum(p["figure"] for p in pages)
    if scanned:
        parts = google_ocr(data, tmp_local, f"{folder}/_ocr_tmp")
        doc = build_docx(stem, data, pages, parts, note=OCR_NOTE)
        how = f"扫描件：OCR 文字 + {figs} 张图页，共 {len(pages)} 页"
    else:
        doc = build_docx(stem, data, pages)
        how = f"文字版：原文文字 + {figs} 张图页，共 {len(pages)} 页"
    out = os.path.join(tmp_local, "doc.docx")
    with open(out, "wb") as fh:
        fh.write(doc)
    upload_as_gdoc(out, ".docx", folder, title)
    return how


# ---------------------------------------------------------------- state & 主流程

def load_state():
    try:
        with open(STATE_PATH, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {"done": {}, "failed": {}}


def save_state(st):
    with open(STATE_PATH, "w", encoding="utf-8") as f:
        json.dump(st, f, ensure_ascii=False, indent=1, sort_keys=True)


def main():
    ap = argparse.ArgumentParser(description="给「湖州线索/<项目>/」里的附件生成 Google 文档版")
    ap.add_argument("--base", default="湖州线索")
    ap.add_argument("--since", type=int, default=7, help="只处理最近几天有公告的项目（0=全部）")
    ap.add_argument("--project", default=None, help="只处理项目名含该关键词的")
    ap.add_argument("--limit", type=int, default=0, help="最多转几个文件（0=不限）")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    rows, _ = load_links()
    projects = build_projects(rows, load_announcements())
    cutoff = (datetime.now(CST) - timedelta(days=a.since)).strftime("%Y-%m-%d") if a.since else ""
    picked = {n: p for n, p in projects.items()
              if (not cutoff or (p["latest"] or "") >= cutoff) and (not a.project or a.project in n)}
    st = load_state()
    todo = [(n, f) for n, p in sorted(picked.items(), key=lambda x: x[1]["latest"] or "", reverse=True)
            for f in p["files"] if f"gdoc|{f['oss']}" not in st["done"]
            and f["filename"].lower().endswith(CONVERTIBLE + ARCHIVE_EXT)]
    log(f"{len(picked)} 个项目（{'全部' if not cutoff else cutoff + ' 以来'}），待转附件 {len(todo)} 个")
    if a.dry_run or not todo:
        for n, f in todo:
            log(f"  {n[:30]} › {f['filename'][:40]}")
        return 0

    converted = failed = 0
    tmp = tempfile.mkdtemp(prefix="huzhou_gdoc_")
    try:
        for i, (name, f) in enumerate(todo, 1):
            if a.limit and converted >= a.limit:
                break
            folder = f"{REMOTE}{sanitize(a.base)}/{sanitize(name)}"
            log(f"[{i}/{len(todo)}] {name[:36]} › {f['filename'][:40]}")
            work = tempfile.mkdtemp(dir=tmp)
            src_key = f"gdoc|{f['oss']}"
            try:
                local = os.path.join(work, "src" + os.path.splitext(f["filename"])[1].lower())
                download(f["oss"], local)
                if f["filename"].lower().endswith(ARCHIVE_EXT):
                    out_dir = os.path.join(work, "x")
                    os.makedirs(out_dir)
                    inner = extract_zip(local, out_dir) if sniff_archive_kind(local) == "zip" else extract_bsdtar(local, out_dir)
                    # to_drive.py 用 rclone copy 上传解出的文件，Drive 上的名字就是解压时的安全名
                    items = [(os.path.join(out_dir, n), n, n) for n, _ in inner]
                else:
                    # to_drive.py 直接上传时，Drive 上的名字是 sanitize(文件名, 120)
                    items = [(local, f["filename"], sanitize(f["filename"], 120))]
                all_ok = True
                for path, fname, raw_name in items:
                    if not fname.lower().endswith(CONVERTIBLE):
                        continue
                    key = f"{src_key}#{fname}"
                    if key in st["done"]:
                        continue
                    try:
                        how = convert_one(path, fname, folder, work)
                        raw_gone = trash_raw(folder, raw_name)
                        st["done"][key] = {"project": name, "file": fname, "how": how, "raw_trashed": raw_gone,
                                           "at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}
                        st["failed"].pop(key, None)
                        converted += 1
                        log(f"        ✓ {fname[:40]}{SUFFIX}：{how}；原文件{'已移到回收站' if raw_gone else '没找到（可能还没同步），未删'}")
                    except Exception as e:  # noqa: BLE001 — 单个失败下次重试
                        all_ok = False
                        failed += 1
                        st["failed"][key] = {"project": name, "file": fname, "error": str(e)[:200]}
                        log(f"        ✗ {fname[:40]}：{str(e)[:160]}")
                    save_state(st)
                if all_ok:
                    st["done"][src_key] = {"project": name, "file": f["filename"],
                                           "at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}
                    save_state(st)
            except Exception as e:  # noqa: BLE001
                failed += 1
                st["failed"][src_key] = {"project": name, "file": f["filename"], "error": str(e)[:200]}
                save_state(st)
                log(f"        ✗ 下载 / 解压失败：{str(e)[:160]}")
            finally:
                shutil.rmtree(work, ignore_errors=True)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    log(f"本次生成文档版 {converted} 份，失败 {failed} 个（失败的下次自动重试）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
